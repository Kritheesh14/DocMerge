import io
import os
import shutil
import tempfile
import uuid
import zipfile
from pathlib import Path

from flask import Flask, jsonify, render_template, request, send_file, session

app = Flask(__name__)
app.secret_key = os.urandom(24)

UPLOAD_FOLDER = tempfile.mkdtemp(prefix="merger_")
ALLOWED_EXTENSIONS = {".pdf", ".pptx", ".ppt"}
MAX_FILES = 20

PT_TO_EMU = 12700


def allowed(filename):
    return Path(filename).suffix.lower() in ALLOWED_EXTENSIONS


def session_dir(sid):
    d = Path(UPLOAD_FOLDER) / sid
    d.mkdir(parents=True, exist_ok=True)
    return d


@app.route("/")
def index():
    if "sid" not in session:
        session["sid"] = str(uuid.uuid4())
    return render_template("index.html")


@app.route("/upload", methods=["POST"])
def upload():
    sid = session.get("sid", str(uuid.uuid4()))
    session["sid"] = sid
    sdir = session_dir(sid)

    order_file = sdir / "order.txt"
    files = order_file.read_text().splitlines() if order_file.exists() else []

    if len(files) >= MAX_FILES:
        return jsonify({"error": f"Maximum {MAX_FILES} files allowed."}), 400

    uploaded = request.files.getlist("files")
    added = []
    for f in uploaded:
        if len(files) >= MAX_FILES:
            break
        if not f.filename or not allowed(f.filename):
            continue
        safe_name = f"{uuid.uuid4().hex}_{Path(f.filename).name}"
        dest = sdir / safe_name
        f.save(dest)
        entry = f"{safe_name}|||{f.filename}"
        files.append(entry)
        added.append({"id": safe_name, "name": f.filename,
                      "ext": Path(f.filename).suffix.lower()})

    order_file.write_text("\n".join(files))
    return jsonify({"added": added, "total": len(files)})


@app.route("/files")
def list_files():
    sid = session.get("sid")
    if not sid:
        return jsonify({"files": []})
    sdir = session_dir(sid)
    order_file = sdir / "order.txt"
    if not order_file.exists():
        return jsonify({"files": []})
    result = []
    for line in order_file.read_text().splitlines():
        if "|||" not in line:
            continue
        safe, orig = line.split("|||", 1)
        result.append({"id": safe, "name": orig,
                       "ext": Path(orig).suffix.lower()})
    return jsonify({"files": result})


@app.route("/reorder", methods=["POST"])
def reorder():
    sid = session.get("sid")
    if not sid:
        return jsonify({"ok": False}), 400
    sdir = session_dir(sid)
    order_file = sdir / "order.txt"
    new_ids = request.json.get("ids", [])
    if not order_file.exists():
        return jsonify({"ok": False}), 400

    mapping = {}
    for line in order_file.read_text().splitlines():
        if "|||" in line:
            safe, orig = line.split("|||", 1)
            mapping[safe] = orig

    new_lines = [f"{nid}|||{mapping[nid]}" for nid in new_ids if nid in mapping]
    order_file.write_text("\n".join(new_lines))
    return jsonify({"ok": True})


@app.route("/remove", methods=["POST"])
def remove():
    sid = session.get("sid")
    if not sid:
        return jsonify({"ok": False}), 400
    sdir = session_dir(sid)
    order_file = sdir / "order.txt"
    file_id = request.json.get("id")
    if not order_file.exists():
        return jsonify({"ok": False}), 400

    lines = order_file.read_text().splitlines()
    new_lines = [l for l in lines if not l.startswith(file_id + "|||")]
    order_file.write_text("\n".join(new_lines))

    target = sdir / file_id
    if target.exists():
        target.unlink()
    return jsonify({"ok": True})


@app.route("/clear", methods=["POST"])
def clear():
    sid = session.get("sid")
    if not sid:
        return jsonify({"ok": True})
    sdir = session_dir(sid)
    if sdir.exists():
        shutil.rmtree(sdir)
    sdir.mkdir(parents=True, exist_ok=True)
    return jsonify({"ok": True})


@app.route("/merge", methods=["POST"])
def merge():
    sid = session.get("sid")
    if not sid:
        return jsonify({"error": "No session"}), 400
    sdir = session_dir(sid)
    order_file = sdir / "order.txt"
    if not order_file.exists():
        return jsonify({"error": "No files uploaded"}), 400

    entries = [(line.split("|||")[0], line.split("|||")[1])
               for line in order_file.read_text().splitlines() if "|||" in line]
    if len(entries) < 2:
        return jsonify({"error": "Upload at least 2 files"}), 400

    req_data = request.get_json(silent=True) or {}
    custom_name = req_data.get("filename", "").strip()
    if not custom_name:
        custom_name = "merged"

    paths = [str(sdir / safe) for safe, _ in entries]
    exts = {Path(orig).suffix.lower() for _, orig in entries}
    all_pdf  = exts <= {".pdf"}
    all_pptx = exts <= {".pptx", ".ppt"}

    mode = (req_data.get("mode") or "auto").lower()
    if mode == "pdf" and not all_pdf:
        return jsonify({"error": "Queue contains presentations. Use Mixed."}), 400
    if mode in ("pptx", "ppt") and not all_pptx:
        return jsonify({"error": "Queue contains PDFs. Use Mixed."}), 400
    use_mixed = mode == "mixed" or (mode == "auto" and not (all_pdf or all_pptx))

    try:
        if not use_mixed and all_pdf:
            out, mime, dl_name = merge_pdfs(paths), "application/pdf", f"{custom_name}.pdf"
        elif not use_mixed and all_pptx:
            out, mime, dl_name = (merge_pptx(paths),
                                  "application/vnd.openxmlformats-officedocument"
                                  ".presentationml.presentation",
                                  f"{custom_name}.pptx")
        else:
            pdf_paths = []
            tmpdir = tempfile.mkdtemp()
            try:
                for safe, orig in entries:
                    p = str(sdir / safe)
                    if Path(orig).suffix.lower() == ".pdf":
                        pdf_paths.append(p)
                    else:
                        pdf_paths.append(pptx_to_pdf(p, tmpdir))
                out = merge_pdfs(pdf_paths)
            finally:
                shutil.rmtree(tmpdir, ignore_errors=True)
            mime, dl_name = "application/pdf", f"{custom_name}.pdf"

        return send_file(out, mimetype=mime,
                         as_attachment=True, download_name=dl_name)
    except Exception as e:
        return jsonify({"error": str(e)}), 500


def pdf_to_docx(pdf_path, output_path):
    from pdf2docx import Converter
    cv = Converter(pdf_path)
    try:
        cv.convert(output_path)
    finally:
        cv.close()
    return output_path


def pdf_to_pptx(pdf_path, output_path):
    import pymupdf
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor

    doc = pymupdf.open(pdf_path)
    try:
        if doc.page_count == 0:
            raise ValueError("The PDF has no pages.")

        prs = Presentation()
        first_page = doc[0]
        prs.slide_width = Emu(int(first_page.rect.width * PT_TO_EMU))
        prs.slide_height = Emu(int(first_page.rect.height * PT_TO_EMU))
        blank_layout = prs.slide_layouts[6]

        for page in doc:
            slide = prs.slides.add_slide(blank_layout)

            for img_info in page.get_image_info(xrefs=True):
                xref = img_info.get("xref", 0)
                bbox = img_info["bbox"]
                if not xref or bbox[2] <= bbox[0] or bbox[3] <= bbox[1]:
                    continue
                try:
                    image_bytes = doc.extract_image(xref)["image"]
                    slide.shapes.add_picture(
                        io.BytesIO(image_bytes),
                        Emu(int(bbox[0] * PT_TO_EMU)),
                        Emu(int(bbox[1] * PT_TO_EMU)),
                        Emu(int((bbox[2] - bbox[0]) * PT_TO_EMU)),
                        Emu(int((bbox[3] - bbox[1]) * PT_TO_EMU)),
                    )
                except Exception:
                    pass

            for block in page.get_text("dict").get("blocks", []):
                if block.get("type") != 0:
                    continue
                bx0, by0, bx1, by1 = block["bbox"]
                if bx1 <= bx0 or by1 <= by0:
                    continue

                textbox = slide.shapes.add_textbox(
                    Emu(int(bx0 * PT_TO_EMU)), Emu(int(by0 * PT_TO_EMU)),
                    Emu(int(max(bx1 - bx0, 10) * PT_TO_EMU)),
                    Emu(int(max(by1 - by0, 10) * PT_TO_EMU)),
                )
                tf = textbox.text_frame
                tf.word_wrap = True
                for margin in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
                    setattr(tf, margin, 0)

                first_para = True
                for line in block.get("lines", []):
                    para = tf.paragraphs[0] if first_para else tf.add_paragraph()
                    first_para = False
                    for span in line.get("spans", []):
                        text = span.get("text", "")
                        if not text:
                            continue
                        run = para.add_run()
                        run.text = text
                        run.font.size = Pt(max(span.get("size", 12), 6))
                        flags = span.get("flags", 0)
                        run.font.bold = bool(flags & (1 << 4))
                        run.font.italic = bool(flags & (1 << 1))
                        font_name = span.get("font", "")
                        if font_name:
                            run.font.name = font_name.split("-")[0].split(",")[0]
                        color_int = span.get("color", 0)
                        try:
                            run.font.color.rgb = RGBColor(
                                (color_int >> 16) & 255,
                                (color_int >> 8) & 255,
                                color_int & 255,
                            )
                        except Exception:
                            pass

        prs.save(output_path)
        return output_path
    finally:
        doc.close()


def merge_pdfs(paths):
    from pypdf import PdfWriter, PdfReader
    writer = PdfWriter()
    for p in paths:
        for page in PdfReader(p).pages:
            writer.add_page(page)
    out = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
    writer.write(out)
    out.close()
    return out.name


def merge_pptx(paths):
    import copy
    from pptx import Presentation
    base = Presentation(paths[0])
    for p in paths[1:]:
        src = Presentation(p)
        for slide in src.slides:
            layout_idx = min(
                len(base.slide_layouts) - 1,
                src.slides.index(slide)
            )
            new_slide = base.slides.add_slide(base.slide_layouts[layout_idx])
            for ph in new_slide.placeholders:
                ph._element.getparent().remove(ph._element)
            for shape in slide.shapes:
                new_slide.shapes._spTree.insert(2, copy.deepcopy(shape._element))
    out = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
    out.close()
    base.save(out.name)
    return out.name


def pptx_to_pdf(pptx_path, out_dir):
    import win32com.client
    import pythoncom
    
    pythoncom.CoInitialize()
    
    abs_pptx_path = os.path.abspath(pptx_path)
    output_pdf_path = os.path.abspath(os.path.join(out_dir, Path(pptx_path).stem + ".pdf"))
    
    powerpoint = None
    presentation = None
    try:
        powerpoint = win32com.client.DispatchEx("PowerPoint.Application")
        presentation = powerpoint.Presentations.Open(abs_pptx_path, WithWindow=False)
        presentation.SaveAs(output_pdf_path, FileFormat=32)
        
    except Exception as e:
        raise RuntimeError(f"Native Windows PowerPoint conversion failed: {str(e)}")
    finally:
        if presentation:
            presentation.Close()
        if powerpoint:
            powerpoint.Quit()
        pythoncom.CoUninitialize()
            
    if not os.path.exists(output_pdf_path):
        raise FileNotFoundError("PowerPoint completed execution, but the output PDF path is missing.")
        
    return output_pdf_path


if __name__ == "__main__":
    app.run(debug=True, port=5000)
